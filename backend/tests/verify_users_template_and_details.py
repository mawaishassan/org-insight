import asyncio
import io
import openpyxl
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.core.database import AsyncSessionLocal
from app.core.models import User, ExternalUser, UserRole
from app.core.security import create_access_token
from sqlalchemy import select

async def main():
    print("=== Verifying System Users Template & User Detail is_external ===")
    async with AsyncSessionLocal() as db:
        # Find org admin
        admin_res = await db.execute(
            select(User).where(User.role == UserRole.ORG_ADMIN, User.organization_id.isnot(None))
        )
        admin = admin_res.scalars().first()
        if not admin:
            print("No org admin found")
            return

        token = create_access_token(subject=admin.id, extra={"role": admin.role.value, "org_id": admin.organization_id})
        headers = {"Authorization": f"Bearer {token}"}

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            # 1. Download standard users template
            res = await client.get(f"/api/users/bulk-template?organization_id={admin.organization_id}", headers=headers)
            assert res.status_code == 200, f"Expected 200, got {res.status_code}"
            wb = openpyxl.load_workbook(io.BytesIO(res.content))
            ws = wb.active
            headers_row = [cell for cell in next(ws.iter_rows(values_only=True))]
            print(f"System Users Template Headers: {headers_row}")
            assert "is_external" not in headers_row, "is_external SHOULD NOT be in system users template!"
            assert headers_row == ["user_name", "password", "unique_user_key", "full_name", "email", "role"], f"Unexpected headers: {headers_row}"

            # Check example rows
            rows = list(ws.iter_rows(values_only=True))
            print(f"Total rows in template: {len(rows)}")
            for r in rows[1:]:
                print(f"Template row: {r}")
                assert len(r) == 6, f"Row length mismatch: {len(r)}"
            print("[PASS] System users template correctly omits is_external column and external rows.")

            # 2. Download external users template
            ext_res = await client.get(f"/api/users/external/template?organization_id={admin.organization_id}", headers=headers)
            assert ext_res.status_code == 200, f"Expected 200, got {ext_res.status_code}"
            ext_wb = openpyxl.load_workbook(io.BytesIO(ext_res.content))
            ext_ws = ext_wb.active
            ext_headers = [cell for cell in next(ext_ws.iter_rows(values_only=True))]
            print(f"External Users Template Headers: {ext_headers}")
            assert ext_headers == ["username", "full_name", "unique_user_key", "description", "is_active"]
            print("[PASS] External users template correctly available as separate file option.")

            # 3. Test GET user detail endpoint for is_external field
            users_res = await client.get(f"/api/users?organization_id={admin.organization_id}", headers=headers)
            assert users_res.status_code == 200
            user_list = users_res.json()
            print(f"Fetched {len(user_list)} users.")

            # Check each user detail
            for u in user_list[:3]:
                detail_res = await client.get(f"/api/users/{u['id']}?organization_id={admin.organization_id}", headers=headers)
                assert detail_res.status_code == 200
                data = detail_res.json()
                assert "is_external" in data, f"is_external field missing in user detail for user {u['id']}"
                print(f"User #{data['id']} ({data['username']}): is_external={data['is_external']}, description={data.get('description')}")

            print("[PASS] User detail endpoint correctly returns is_external.")

    print("=== ALL CHECKS PASSED SUCCESSFULLY ===")

if __name__ == "__main__":
    asyncio.run(main())
